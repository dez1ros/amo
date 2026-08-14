from copy import copy
from decimal import Decimal, InvalidOperation
from math import ceil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.cell_range import MultiCellRange


ITEM_START_ROW = 20
TEMPLATE_ITEM_END_ROW = 23
TEMPLATE_DYNAMIC_END_ROW = 24
TEMPLATE_DYNAMIC_ROW_COUNT = TEMPLATE_DYNAMIC_END_ROW - ITEM_START_ROW + 1
ITEM_STYLE_TEMPLATE_ROW = 21
TEMPLATE_TOTAL_ROW = 25
TEMPLATE_FOUNDATION_ROW = 33
TEMPLATE_PRINT_END_ROW = 50
ITEM_COLUMNS = ("Y", "BD", "BM", "BV", "CG", "CV", "DC", "DJ", "DW")


def _has_value(value):
    return value is not None and (not isinstance(value, str) or value.strip() != "")


def _to_decimal(value, field_name):
    if not _has_value(value):
        return None

    normalized = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return Decimal(normalized)
    except (InvalidOperation, ValueError) as error:
        raise ValueError(
            f"Некорректное числовое значение для {field_name}: {value!r}"
        ) from error


def _to_excel_number(value, field_name):
    number = _to_decimal(value, field_name)
    if number is None:
        return None
    if number == number.to_integral_value():
        return int(number)
    return number


def _item_text(item):
    parts = []
    for key in ("name", "desc"):
        value = item.get(key)
        if _has_value(value):
            parts.append(str(value).strip())
    return "\n".join(parts)


def _item_row_height(text):
    visual_lines = sum(max(1, ceil(len(line) / 30)) for line in text.splitlines())
    return max(15, visual_lines * 15)


def _shifted_row(row, item_count):
    if row < ITEM_START_ROW:
        return row
    if row <= TEMPLATE_DYNAMIC_END_ROW:
        return row if row < ITEM_START_ROW + item_count else None
    return row + item_count - TEMPLATE_DYNAMIC_ROW_COUNT


def _resize_item_rows(sheet, item_count):
    row_delta = item_count - TEMPLATE_DYNAMIC_ROW_COUNT
    original_merges = [
        (cell_range.min_col, cell_range.min_row, cell_range.max_col, cell_range.max_row)
        for cell_range in sheet.merged_cells.ranges
    ]
    item_merge_columns = [
        (min_col, max_col)
        for min_col, min_row, max_col, max_row in original_merges
        if min_row == ITEM_STYLE_TEMPLATE_ROW and max_row == ITEM_STYLE_TEMPLATE_ROW
    ]
    template_row_styles = {
        column: copy(sheet.cell(ITEM_STYLE_TEMPLATE_ROW, column)._style)
        for column in range(1, sheet.max_column + 1)
    }
    template_row_dimension = copy(sheet.row_dimensions[ITEM_STYLE_TEMPLATE_ROW])
    original_row_dimensions = {
        row: copy(dimension) for row, dimension in sheet.row_dimensions.items()
    }

    if row_delta > 0:
        sheet.insert_rows(TEMPLATE_TOTAL_ROW, amount=row_delta)
    elif row_delta < 0:
        sheet.delete_rows(ITEM_START_ROW + item_count, amount=-row_delta)

    sheet.merged_cells = MultiCellRange()
    for coordinate, cell in list(sheet._cells.items()):
        if isinstance(cell, MergedCell):
            del sheet._cells[coordinate]

    shifted_merges = []
    for min_col, min_row, max_col, max_row in original_merges:
        if max_row < ITEM_START_ROW:
            shifted_merges.append((min_col, min_row, max_col, max_row))
        elif min_row >= TEMPLATE_TOTAL_ROW:
            shifted_merges.append(
                (min_col, min_row + row_delta, max_col, max_row + row_delta)
            )
        elif min_row >= ITEM_START_ROW and max_row <= TEMPLATE_DYNAMIC_END_ROW:
            if (
                max_row < ITEM_START_ROW + item_count
                and min_row <= TEMPLATE_ITEM_END_ROW
            ):
                shifted_merges.append((min_col, min_row, max_col, max_row))
        else:
            shifted_merges.append((min_col, min_row, max_col, max_row + row_delta))

    if item_count > TEMPLATE_ITEM_END_ROW - ITEM_START_ROW + 1:
        for row in range(TEMPLATE_ITEM_END_ROW + 1, ITEM_START_ROW + item_count):
            for column, style in template_row_styles.items():
                sheet.cell(row, column)._style = copy(style)
            for min_col, max_col in item_merge_columns:
                shifted_merges.append((min_col, row, max_col, row))

    sheet.row_dimensions.clear()
    for old_row, dimension in original_row_dimensions.items():
        new_row = _shifted_row(old_row, item_count)
        if new_row is not None:
            dimension.index = new_row
            sheet.row_dimensions[new_row] = dimension
    if item_count > TEMPLATE_ITEM_END_ROW - ITEM_START_ROW + 1:
        for row in range(TEMPLATE_ITEM_END_ROW + 1, ITEM_START_ROW + item_count):
            dimension = copy(template_row_dimension)
            dimension.index = row
            sheet.row_dimensions[row] = dimension

    for min_col, min_row, max_col, max_row in shifted_merges:
        sheet.merge_cells(
            start_row=min_row,
            start_column=min_col,
            end_row=max_row,
            end_column=max_col,
        )

    sheet.print_area = f"A1:HK{TEMPLATE_PRINT_END_ROW + row_delta}"
    return row_delta


def generate_upd_xlsx(template_path, output_path, context):
    template_path = Path(template_path)
    output_path = Path(output_path)
    items = context.get("items") or []

    if not template_path.is_file():
        raise FileNotFoundError(f"Шаблон УПД не найден: {template_path}")
    if template_path.resolve() == output_path.resolve():
        raise ValueError("Выходной файл УПД не должен перезаписывать шаблон")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(template_path)
    try:
        sheet = workbook.active
        row_delta = _resize_item_rows(sheet, len(items))
        item_rows = range(ITEM_START_ROW, ITEM_START_ROW + len(items))
        total_row = TEMPLATE_TOTAL_ROW + row_delta
        foundation_row = TEMPLATE_FOUNDATION_ROW + row_delta

        number = context.get("number")
        date = context.get("date")
        client = context.get("client")
        phone = context.get("phone")

        if _has_value(number):
            sheet["AM1"] = number
        if _has_value(date):
            sheet["BC1"] = date
        if _has_value(client):
            recipient = str(client).strip()
            if _has_value(phone):
                recipient = f"{recipient}, {str(phone).strip()}"
            sheet["AZ8"] = recipient
            sheet["AZ11"] = client
        if _has_value(number) and _has_value(date):
            sheet[f"AR{foundation_row}"] = f"Счёт-Договор № {number} от {date} г."

        for row in item_rows:
            for column in ITEM_COLUMNS:
                sheet[f"{column}{row}"] = None

        for row, item in zip(item_rows, items):
            item_text = _item_text(item)
            sheet[f"Y{row}"] = item_text
            sheet.row_dimensions[row].height = _item_row_height(item_text)
            sheet[f"BD{row}"] = "шт."
            if _has_value(item.get("qty")):
                sheet[f"BM{row}"] = _to_excel_number(item["qty"], "qty")
            if _has_value(item.get("price")):
                sheet[f"BV{row}"] = _to_excel_number(item["price"], "price")
            if _has_value(item.get("total")):
                item_total = _to_excel_number(item["total"], "total")
                sheet[f"CG{row}"] = item_total
                sheet[f"DW{row}"] = item_total
            sheet[f"CV{row}"] = "Без акциза"
            sheet[f"DC{row}"] = "Без НДС"
            sheet[f"DJ{row}"] = "Без НДС"

        if _has_value(context.get("total_sum")):
            total_sum = _to_excel_number(context["total_sum"], "total_sum")
            sheet[f"CG{total_row}"] = total_sum
            sheet[f"DW{total_row}"] = total_sum

        workbook.save(output_path)
    finally:
        workbook.close()
