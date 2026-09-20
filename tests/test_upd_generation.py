import hashlib
import io
import json
import tempfile
import unittest
import uuid
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docxtpl import DocxTemplate
from openpyxl import load_workbook

from generation import ITEM_COLUMNS, generate_upd_xlsx
from utils import add_file_to_archive


PROJECT_ROOT = Path(__file__).resolve().parents[1]
UPD_TEMPLATE = PROJECT_ROOT / "docs" / "template_upd_veng.xlsx"


def file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def sample_context(item_count=1):
    items = [
        {
            "name": f"Товар {index + 1}",
            "desc": "Описание",
            "qty": index + 1,
            "price": "1250,50",
            "total": "1250.50",
        }
        for index in range(item_count)
    ]
    return {
        "production_dates": "15-35",
        "number": "42-26",
        "client": "ООО Покупатель",
        "deal": "42-26",
        "phone": "+7 900 000-00-00",
        "organization": "proh",
        "date": "14.08.2026",
        "date_words": "14 августа 2026",
        "items": items,
        "total_sum": "1250,50",
        "total_words": "Одна тысяча двести пятьдесят",
        "half_sum": "625.25",
        "half_words": "Шестьсот двадцать пять",
        "date_day": "14",
        "date_month": "08",
        "date_year": 2026,
    }


class UpdGenerationTests(unittest.TestCase):
    def test_self_employed_templates_have_only_new_organization_details(self):
        docx_templates = sorted(
            (PROJECT_ROOT / "docs").glob("template_*_veng_self.docx")
        )
        self.assertEqual(6, len(docx_templates))

        for template_path in docx_templates:
            with self.subTest(template=template_path.name):
                with zipfile.ZipFile(template_path) as archive:
                    document_xml = archive.read("word/document.xml")
                root = ET.fromstring(document_xml)
                text = "".join(
                    node.text or "" for node in root.iter() if node.tag.endswith("}t")
                )

                self.assertIn("Самозанятый Венгеров А.Ю.", text)
                self.assertIn("770200924557", text)
                self.assertIn("+79772730883", text)
                self.assertNotIn("Венгерова Е.Ю.", text)
                self.assertNotIn("772573133488", text)

        upd_template = PROJECT_ROOT / "docs" / "template_upd_veng_self.xlsx"
        upd_book = load_workbook(upd_template)
        try:
            sheet = upd_book.active
            self.assertEqual("Самозанятый Венгеров А.Ю.", sheet["AZ4"].value)
            self.assertEqual("ИНН 770200924557", sheet["AZ6"].value)
            self.assertEqual("Венгеров А.Ю.", sheet["BN27"].value)
            self.assertIsNone(sheet["AZ5"].value)
        finally:
            upd_book.close()

    def test_generates_one_item_row_without_changing_template(self):
        template_hash = file_hash(UPD_TEMPLATE)

        template_book = load_workbook(UPD_TEMPLATE)
        template_sheet = template_book.active
        original_styles = {
            coordinate: tuple(template_sheet[coordinate]._style)
            for coordinate in ("AM1", "BC1", "AZ8", "AZ11", "Y20", "BM20")
        }
        original_total_styles = {
            coordinate: tuple(template_sheet[coordinate]._style)
            for coordinate in ("CG25", "DW25", "AR33")
        }
        template_book.close()

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "nested" / "УПД 42-26.xlsx"
            generate_upd_xlsx(UPD_TEMPLATE, output_path, sample_context())

            self.assertTrue(output_path.is_file())
            self.assertEqual(template_hash, file_hash(UPD_TEMPLATE))

            output_book = load_workbook(output_path)
            output_sheet = output_book.active
            self.assertEqual("42-26", output_sheet["AM1"].value)
            self.assertEqual("14.08.2026", output_sheet["BC1"].value)
            self.assertEqual(
                "ООО Покупатель, +7 900 000-00-00", output_sheet["AZ8"].value
            )
            self.assertEqual("ООО Покупатель", output_sheet["AZ11"].value)
            self.assertEqual("Товар 1\nОписание", output_sheet["Y20"].value)
            self.assertEqual(30, output_sheet.row_dimensions[20].height)
            self.assertEqual(1, output_sheet["BM20"].value)
            self.assertEqual(1250.5, output_sheet["BV20"].value)
            self.assertEqual(1250.5, output_sheet["CG20"].value)
            self.assertEqual(1250.5, output_sheet["DW21"].value)
            self.assertEqual(
                "Счёт-Договор № 42-26 от 14.08.2026 г.", output_sheet["AR29"].value
            )
            for coordinate in ("BM20", "BV20", "CG20", "DW20", "CG21", "DW21"):
                self.assertEqual("n", output_sheet[coordinate].data_type)
            self.assertEqual(
                original_styles,
                {
                    coordinate: tuple(output_sheet[coordinate]._style)
                    for coordinate in original_styles
                },
            )
            self.assertEqual(
                original_total_styles["CG25"], tuple(output_sheet["CG21"]._style)
            )
            self.assertEqual(
                original_total_styles["DW25"], tuple(output_sheet["DW21"]._style)
            )
            self.assertEqual(
                original_total_styles["AR33"], tuple(output_sheet["AR29"]._style)
            )
            output_merges = {
                str(cell_range) for cell_range in output_sheet.merged_cells.ranges
            }
            self.assertIn("Y20:AV20", output_merges)
            self.assertNotIn("Y21:AV21", output_merges)
            self.assertIn("U21:CF21", output_merges)
            self.assertEqual(46, output_sheet.max_row)
            self.assertEqual("'стр.1'!$A$1:$HK$46", str(output_sheet.print_area))
            output_book.close()

    def test_expands_table_to_seven_item_rows(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "upd.xlsx"
            generate_upd_xlsx(UPD_TEMPLATE, output_path, sample_context(7))

            output_book = load_workbook(output_path)
            output_sheet = output_book.active
            output_merges = {
                str(cell_range) for cell_range in output_sheet.merged_cells.ranges
            }
            for index, row in enumerate(range(20, 27), start=1):
                self.assertEqual(f"Товар {index}\nОписание", output_sheet[f"Y{row}"].value)
                self.assertIn(f"Y{row}:AV{row}", output_merges)
                self.assertEqual(30, output_sheet.row_dimensions[row].height)
            self.assertEqual(
                tuple(output_sheet["Y21"]._style), tuple(output_sheet["Y24"]._style)
            )
            self.assertNotIn("Y27:AV27", output_merges)
            self.assertIn("T24:X24", output_merges)
            self.assertEqual(1250.5, output_sheet["DW27"].value)
            self.assertEqual(
                "Счёт-Договор № 42-26 от 14.08.2026 г.", output_sheet["AR35"].value
            )
            self.assertEqual(52, output_sheet.max_row)
            self.assertEqual("'стр.1'!$A$1:$HK$52", str(output_sheet.print_area))
            output_book.close()

    def test_adds_upd_to_zip_without_renaming_existing_entries(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_dir = Path(temp_dir)
            upd_path = temp_dir / "upd.xlsx"
            archive_path = temp_dir / "documents.zip"
            generate_upd_xlsx(UPD_TEMPLATE, upd_path, sample_context())

            with zipfile.ZipFile(archive_path, "w") as archive:
                archive.writestr("Существующий документ.docx", b"docx")
                add_file_to_archive(archive, upd_path, "УПД 42-26.xlsx")

            with zipfile.ZipFile(archive_path) as archive:
                self.assertEqual(
                    ["Существующий документ.docx", "УПД 42-26.xlsx"],
                    archive.namelist(),
                )
                self.assertEqual(b"docx", archive.read("Существующий документ.docx"))

    def test_existing_docx_template_still_renders(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "act.docx"
            document = DocxTemplate(PROJECT_ROOT / "docs" / "template_act_proh.docx")
            document.render(sample_context())
            document.save(output_path)

            self.assertTrue(output_path.is_file())
            Document(output_path)

    def test_flask_selects_matching_upd_template(self):
        from main import app

        cases = (
            ("proh", "Прохоров М.В.", "template_upd_proh.xlsx"),
            ("veng", "Венгерова Е.Ю.", "template_upd_veng.xlsx"),
            ("veng_self", "Венгеров А.Ю.", "template_upd_veng_self.xlsx"),
        )
        for organization, expected_seller, template_name in cases:
            with self.subTest(organization=organization):
                number = f"codex-upd-{uuid.uuid4().hex}"
                generated_zip = PROJECT_ROOT / "generated" / f"{number}.zip"
                saved_form = PROJECT_ROOT / "saved_forms" / f"{number}.json"
                response = None
                template_path = PROJECT_ROOT / "docs" / template_name
                template_hash = file_hash(template_path)
                existing_names = [
                    f"Счёт-договор {number}.docx",
                    f"Акт {number}.docx",
                    f"Накладная {number}.docx",
                    f"Спецификация {number}.docx",
                    f"Счёт предоплата к {number}.docx",
                    f"Счёт остаток {number}.docx",
                ]
                if organization == "proh":
                    existing_names.insert(1, f"Счёт-договор {number} без QR.docx")
                    existing_names.extend(
                        ["КП.docx", f"Кассовый ордер {number}.docx"]
                    )
                existing_names.append(f"Все документы {number}.docx")

                request_data = {
                        "client": "ООО Покупатель",
                        "deal": number,
                        "phone": "+7 900 000-00-00",
                        "production_dates": "15-35",
                        "organization": organization,
                        "name[]": ["Тестовый товар"],
                        "desc[]": ["Описание"],
                        "qty[]": ["2"],
                        "price[]": ["100.50"],
                }
                try:
                    with app.test_client() as client:
                        response = client.post("/", data=request_data)

                    self.assertEqual(200, response.status_code)
                    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
                        self.assertEqual(existing_names, archive.namelist()[:-1])
                        upd_name = f"УПД {number}.xlsx"
                        self.assertEqual(upd_name, archive.namelist()[-1])
                        upd_book = load_workbook(io.BytesIO(archive.read(upd_name)))
                        self.assertEqual(expected_seller, upd_book.active["BN23"].value)
                        upd_book.close()
                    saved_data = json.loads(saved_form.read_text(encoding="utf-8"))
                    self.assertEqual(organization, saved_data["organization"])
                    self.assertNotIn("is_ip", saved_data)
                    self.assertEqual(template_hash, file_hash(template_path))
                finally:
                    if response is not None:
                        response.close()
                    generated_zip.unlink(missing_ok=True)
                    saved_form.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
